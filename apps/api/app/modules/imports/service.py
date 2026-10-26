from app.modules.imports.jobs import process as process
from app.modules.imports import repository
from app.modules.verification.matching import make_indexes as make_indexes
from app.modules.verification.rules import evaluate as evaluate
from app.modules.imports.mapping import ALIASES as ALIASES, MANDATORY as MANDATORY, normalized as normalized
from datetime import datetime, date
from io import BytesIO, StringIO
import csv
from uuid import uuid4
import zipfile
from openpyxl import load_workbook
from psycopg.types.json import Jsonb
from app.core.errors import fail
from app.core.files import store
from app.db.session import connect, many
from app.modules.students.service import jsonable
from app.core.http import listing, pagination, batch_scope, idem
import hashlib
from uuid import UUID
from fastapi import Request, UploadFile, File, Form, BackgroundTasks
from fastapi.responses import Response
from app.core.security import allow, institution_scope, reason
from app.core.audit import append
from app.core.files import contained, root_path
from app.settings import settings
from app.modules.imports import service as imports
from app.core.contracts import Reason
from app.modules.imports.schemas import Mapping
from app.modules.imports.schemas import Confirm


def parse(data, filename):
    if len(data) > 20 * 1024 * 1024:
        fail("FILE_TOO_LARGE", 413)
    if filename.lower().endswith(".xlsx"):
        try:
            archive = zipfile.ZipFile(BytesIO(data))
            if sum((i.file_size for i in archive.infolist())) > 100 * 1024 * 1024:
                fail("XLSX_EXPANSION_LIMIT")
            book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
            nonempty = []
            for sheet in book:
                if sheet.max_column and sheet.max_column > 100:
                    fail("INVALID_HEADERS")
                if sheet.max_row and sheet.max_row > 1001:
                    fail("TOO_MANY_ROWS")
                rows = list(sheet.iter_rows())
                if any((any((c.value is not None for c in row)) for row in rows)):
                    nonempty.append(rows)
            if len(nonempty) != 1:
                fail("WORKSHEET_SELECTION_REQUIRED")
            cells = nonempty[0]
            if any((c.data_type == "f" for row in cells for c in row)):
                fail("XLSX_FORMULA_NOT_SUPPORTED")
            headers = [str(c.value or "") for c in cells[0]]
            rows = [
                {
                    h: c.value.date().isoformat()
                    if isinstance(c.value, datetime)
                    else c.value.isoformat()
                    if isinstance(c.value, date)
                    else str(c.value)
                    if c.value is not None
                    else ""
                    for h, c in zip(headers, row)
                }
                for row in cells[1:]
            ]
        except (zipfile.BadZipFile, ValueError):
            fail("UNSUPPORTED_SPREADSHEET")
    elif filename.lower().endswith(".csv"):
        try:
            reader = csv.DictReader(StringIO(data.decode("utf-8-sig")))
            headers = list(reader.fieldnames or [])
            rows = list(reader)
        except (UnicodeDecodeError, csv.Error):
            fail("UNSUPPORTED_CSV")
    else:
        fail("UNSUPPORTED_MEDIA_TYPE", 415)
    if len(headers) > 100 or len(set(headers)) != len(headers):
        fail("INVALID_HEADERS")
    if not rows:
        fail("EMPTY_FILE")
    if len(rows) > 1000:
        fail("TOO_MANY_ROWS")
    return (headers, rows)


def upload(conn, user, data, filename, key):
    headers, rows = parse(data, filename)
    batch = uuid4()
    storage, sha = store(
        data, user["institution_id"], "imports", "xlsx" if filename.lower().endswith(".xlsx") else "csv"
    )
    repository.upload_import_batches(
        conn, (batch, user["institution_id"], user["id"], filename, storage, sha, len(rows))
    )
    for i, row in enumerate(rows, 1):
        repository.upload_import_rows(conn, (batch, i, Jsonb(row)))
    return dict(
        batch_id=batch,
        status="UPLOADED",
        total_rows=len(rows),
        headers=headers,
        preview=[dict(row_no=i + 1, values=row) for i, row in enumerate(rows[:20])],
        suggested_mapping={k: dict(column=v, confidence=100) for k, v in ALIASES.items() if v in headers},
    )


def start(conn, batch):
    refs = repository.start_students(conn, (batch["institution_id"],))
    claims = {r["apaar_id"]: str(r["id"]) for r in repository.start_students_2(conn, ())}
    snapshot = dict(rows=jsonable(refs), claims=claims, rule_version=1, demo_as_of="2026-10-25")
    repository.start_import_batches(conn, (Jsonb(snapshot), batch["id"]))
    job = uuid4()
    repository.start_import_jobs(conn, (job, batch["id"]))
    return job


def report(conn, batch_id):
    batch = repository.report_import_batches(conn, (batch_id,))
    buckets = repository.report_import_errors(conn, (batch_id,))
    all_count = repository.report_import_errors_2(conn, (batch_id,))["n"]
    return dict(
        batch_id=batch_id,
        clean=batch["clean_rows"],
        flagged=batch["flagged_rows"],
        incomplete=batch["incomplete_rows"],
        affected_rows=batch["flagged_rows"] + batch["incomplete_rows"],
        primary_findings=buckets,
        all_findings=all_count,
        secondary_findings=all_count - sum((b["count"] for b in buckets)),
        processed_rows=batch["processed_rows"],
        total_rows=batch["total_rows"],
        status=batch["status"],
    )


def upload_import(request: Request, file: UploadFile = File(...), institution_id: UUID = Form(...), user=None):
    allow(user, "OPERATOR")
    institution_scope(user, institution_id)
    data = file.file.read(20 * 1024 * 1024 + 1)
    with connect() as conn:
        return idem(
            conn,
            user,
            request,
            dict(sha=hashlib.sha256(data).hexdigest(), institution=str(institution_id)),
            lambda key: imports.upload(conn, user, data, file.filename or "", key),
        )


def list_imports(page: int = 1, page_size: int = 25, user=None):
    allow(user, "OPERATOR", "OFFICER", "ADMIN")
    offset = pagination(page, page_size)
    ids = institution_scope(user, admin=True)
    with connect() as conn:
        cols = (
            "id,institution_id,status,total_rows,processed_rows,clean_rows,flagged_rows,incomplete_rows,created_at"
            + ("" if user["role_code"] == "ADMIN" else ",filename")
        )
        rows = many(
            conn,
            f"SELECT {cols} FROM import_batches WHERE (%s::uuid[] IS NULL OR institution_id=ANY(%s::uuid[])) ORDER BY created_at DESC LIMIT %s OFFSET %s",
            (ids, ids, page_size, offset),
        )
        total = repository.list_imports_import_batches(conn, (ids, ids))["n"]
        return listing(rows, total, page, page_size)


def import_detail(batch_id: UUID, user=None):
    with connect() as conn:
        batch = batch_scope(conn, user, batch_id)
        rows = repository.import_detail_import_rows(conn, (batch_id,))
        job = repository.import_detail_import_jobs(conn, (batch_id,))
        return {
            **{k: batch[k] for k in ["id", "filename", "status", "total_rows", "processed_rows", "column_mapping"]},
            "preview": rows,
            "report": imports.report(conn, batch_id),
            "job": job,
        }


def set_mapping(batch_id: UUID, body: Mapping, user=None):
    allow(user, "OPERATOR")
    with connect() as conn:
        batch = batch_scope(conn, user, batch_id)
        if batch["status"] not in {"UPLOADED", "MAPPED"}:
            fail("IMPORT_ALREADY_STARTED", 409)
        rows = repository.set_mapping_import_rows(conn, (batch_id,))
        headers = set(rows[0]["raw_data"])
        if set(body.mapping.values()) - headers or set(body.mapping) - set(imports.ALIASES):
            fail("UNKNOWN_COLUMN", 400)
        if len(set(body.mapping.values())) != len(body.mapping):
            fail("DUPLICATE_MAPPING", 400)
        if imports.MANDATORY - set(body.mapping):
            fail("REQUIRED_MAPPING_MISSING")
        if body.date_format not in {"DD/MM/YYYY", "ISO"}:
            fail("DATE_FORMAT_UNSUPPORTED")
        repository.set_mapping_import_batches(conn, (Jsonb(body.model_dump()), batch_id))
        return dict(
            batch_id=batch_id,
            status="MAPPED",
            required_fields_mapped=True,
            preview=[
                dict(
                    row_no=r["row_no"],
                    normalized=imports.normalized(r["raw_data"], body.mapping, body.date_format),
                    errors=[],
                )
                for r in rows
            ],
        )


def validate_import(batch_id: UUID, body: Confirm, request: Request, tasks: BackgroundTasks, user=None):
    allow(user, "OPERATOR")
    if not body.confirm_synthetic_data:
        fail("SYNTHETIC_CONFIRMATION_REQUIRED")
    with connect() as conn:
        batch = batch_scope(conn, user, batch_id)

        def work(key):
            if batch["status"] != "MAPPED":
                fail("MAPPING_REQUIRED", 409)
            job = imports.start(conn, batch)
            tasks.add_task(imports.process, job)
            return dict(job_id=job, batch_id=batch_id, status="QUEUED", poll_url=f"/api/v1/import-jobs/{job}")

        result = idem(conn, user, request, body.model_dump(), work)
    return result


def get_job(job_id: UUID, user=None):
    with connect() as conn:
        job = repository.get_job_import_jobs(conn, (job_id,))
        if not job:
            fail("NOT_FOUND", 404)
        batch = batch_scope(conn, user, job["batch_id"])
        return {**job, "processed_rows": batch["processed_rows"], "total_rows": batch["total_rows"]}


def retry_job(job_id: UUID, body: Reason, request: Request, tasks: BackgroundTasks, user=None):
    allow(user, "OPERATOR")
    reason(body.reason)
    with connect() as conn:
        job = repository.retry_job_import_jobs(conn, (job_id,))
        if not job:
            fail("NOT_FOUND", 404)
        batch = batch_scope(conn, user, job["batch_id"])

        def work(key):
            if job["status"] != "FAILED":
                fail("JOB_NOT_FAILED", 409)
            if job["retry_count"] >= 3:
                fail("RETRY_LIMIT", 409)
            if not contained(root_path(settings.storage_root), batch["storage_key"]).is_file():
                fail("UPLOAD_MISSING", 409)
            repository.retry_job_import_jobs_2(conn, (job_id,))
            repository.queue_batch(conn, (job["batch_id"],))
            append(
                conn,
                batch["institution_id"],
                user["id"],
                "IMPORT_RETRIED",
                "import",
                batch["id"],
                body.reason,
                dict(next_row=job["next_row"]),
            )
            tasks.add_task(imports.process, job_id)
            return dict(
                id=job_id,
                job_id=job_id,
                status="QUEUED",
                retry_count=job["retry_count"] + 1,
                resume_from_row=job["next_row"],
            )

        return idem(conn, user, request, body.model_dump(), work)


def import_report(batch_id: UUID, user=None):
    with connect() as conn:
        batch_scope(conn, user, batch_id)
        return imports.report(conn, batch_id)


def errors_csv(batch_id: UUID, user=None):
    with connect() as conn:
        batch_scope(conn, user, batch_id)
        rows = repository.errors_csv_import_errors(conn, (batch_id,))
    output = StringIO()
    writer = csv.writer(output)
    writer.writerow(["row_no", "field", "code", "is_primary", "message", "submitted_value", "expected"])

    def safe(value):
        value = str(value or "")
        return "'" + value if value.startswith(("=", "+", "-", "@", "\t", "\r")) else value

    for row in rows:
        writer.writerow(
            [safe(row[k]) for k in ["row_no", "field", "code", "is_primary", "message"]]
            + [safe(row["evidence"].get(k)) for k in ["submitted_value", "expected"]]
        )
    return Response(
        output.getvalue(),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="import-errors.csv"'},
    )
